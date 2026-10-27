#!/usr/bin/env python3
"""Generate Word document with RCLS snail row feasibility tables."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parent.parent / "Markdown History" / "RCLS-Snail-Feasible-Rows-Summary.docx"

ROW1 = [
    [1, 3, 4, 6, 2, 9, 5, 7, 8, 10],
    [1, 3, 5, 9, 4, 7, 2, 6, 8, 10],
    [1, 3, 6, 2, 4, 7, 9, 5, 8, 10],
    [1, 3, 7, 5, 2, 9, 6, 4, 8, 10],
    [1, 5, 4, 2, 3, 8, 9, 7, 6, 10],
    [1, 5, 4, 9, 3, 8, 2, 7, 6, 10],
    [1, 5, 7, 2, 3, 8, 9, 4, 6, 10],
    [1, 5, 7, 9, 3, 8, 2, 4, 6, 10],
    [1, 6, 4, 2, 3, 8, 9, 7, 5, 10],
    [1, 6, 4, 9, 3, 8, 2, 7, 5, 10],
    [1, 6, 7, 2, 3, 8, 9, 4, 5, 10],
    [1, 6, 7, 9, 3, 8, 2, 4, 5, 10],
    [1, 8, 4, 5, 2, 9, 6, 7, 3, 10],
    [1, 8, 5, 2, 4, 7, 9, 6, 3, 10],
    [1, 8, 6, 9, 4, 7, 2, 5, 3, 10],
    [1, 8, 7, 6, 2, 9, 5, 4, 3, 10],
]

ROW2 = [
    [1, 8, 4, 7, 3, 10],
    [1, 8, 7, 4, 3, 10],
    [4, 5, 10, 1, 6, 7],
    [4, 6, 1, 10, 5, 7],
    [5, 1, 7, 4, 10, 6],
    [5, 10, 4, 7, 1, 6],
    [6, 1, 4, 7, 10, 5],
    [6, 10, 7, 4, 1, 5],
    [7, 5, 1, 10, 6, 4],
    [7, 6, 10, 1, 5, 4],
    [10, 8, 4, 7, 3, 1],
    [10, 8, 7, 4, 3, 1],
]

ROW3 = [
    [1, 2, 9, 10],
    [1, 9, 2, 10],
    [2, 1, 10, 9],
    [2, 10, 1, 9],
    [5, 2, 9, 6],
    [5, 9, 2, 6],
    [6, 2, 9, 5],
    [6, 9, 2, 5],
    [9, 1, 10, 2],
    [9, 10, 1, 2],
    [10, 2, 9, 1],
    [10, 9, 2, 1],
]

ROW4 = [
    [1, 10],
    [3, 8],
    [8, 3],
    [10, 1],
]


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def style_header_row(table, fill: str = "2F5496") -> None:
    for cell in table.rows[0].cells:
        set_cell_shading(cell, fill)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
                run.font.name = "Calibri"


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list],
    col_widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for j, h in enumerate(headers):
        hdr[j].text = str(h)
        hdr[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_header_row(table)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = str(val)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
        if i % 2 == 0:
            for cell in table.rows[i].cells:
                set_cell_shading(cell, "D6E4F0")
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)


def main() -> None:
    doc = Document()
    title = doc.add_heading("RCLS Snail Feasibility Summary", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Order n = 10 · Diagonal-only starting grid (upper-V)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(12)
    sub.runs[0].italic = True

    doc.add_heading("Setup", level=1)
    for line in [
        "Givens: 20 cells on the main and anti-diagonals (X-pattern only).",
        "RCLS: each row and column is a permutation of 1–10; all 90 horizontal ordered pairs (x→y, x≠y) are distinct.",
        "Snail: cells linked on each concentric shell (snail method) must share the same value.",
        "Feasibility: OR-Tools CP-SAT — does a full completion exist?",
        "Rows 1–4 were analyzed separately; a feasible Row 2 pattern is not guaranteed to combine with a feasible Row 3 pattern.",
    ]:
        p = doc.add_paragraph(line, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)

    doc.add_heading("Row 1 — Complete top row", level=1)
    doc.add_paragraph(
        "Fixed: column 1 = 1, column 10 = 10. "
        "4,752 column-feasible candidates → 16 admit snail + RCLS."
    )
    h1 = ["#"] + [f"C{j}" for j in range(1, 11)]
    add_table(doc, h1, [[i + 1] + r for i, r in enumerate(ROW1)])

    doc.add_page_break()
    doc.add_heading("Row 2 — Columns 3–8 (between 2 and 9)", level=1)
    doc.add_paragraph(
        "Fixed on row 2: C2 = 2, C9 = 9. Columns 1 and 10 left blank in the test. "
        "Pattern: _  2  ···  9  _.  3,800 filtered candidates → 12 feasible."
    )
    add_table(doc, ["#", "C3", "C4", "C5", "C6", "C7", "C8"], [[i + 1] + r for i, r in enumerate(ROW2)])

    doc.add_heading("Row 3 — Columns 4–7 (between 3 and 8)", level=1)
    doc.add_paragraph(
        "Fixed on row 3: C3 = 3, C8 = 8. Row 2 not considered. "
        "Pattern: _  _  3  ···  8  _  _.  524 candidates → 12 feasible."
    )
    add_table(doc, ["#", "C4", "C5", "C6", "C7"], [[i + 1] + r for i, r in enumerate(ROW3)])

    doc.add_heading("Row 4 — Columns 5–6 (between 4 and 7)", level=1)
    doc.add_paragraph(
        "Fixed on row 4: C4 = 4, C7 = 7. Rows 2–3 not considered. "
        "Pattern: _  _  _  4  ···  7  _  _  _.  30 candidates → 4 feasible."
    )
    add_table(doc, ["#", "C5", "C6"], [[i + 1] + r for i, r in enumerate(ROW4)])

    doc.add_paragraph()
    foot = doc.add_paragraph(
        "Generated from RCLS-Validation-Shiny enumeration (scripts/top_row_snail_enum.py, "
        "scripts/row_segment_snail_enum.py)."
    )
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
